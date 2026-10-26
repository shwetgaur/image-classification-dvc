#!/usr/bin/env python3
"""Generate AWS deployment assignment report (Word)."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

OUT = Path(r"C:\Users\shwet\Downloads\AWS_DVC_Docker_Deployment_Report.docx")
ROOT = Path(__file__).resolve().parents[1]
REPO = "https://github.com/shwetgaur/image-classification-dvc"
EC2_IP_FILE = ROOT / "deploy" / "ec2_public_ip.txt"

ACCOUNT_ID = "746914061275"
ACCOUNT_NAME = "shwetgaur"
BUCKET = "image-classification-dvc-746914061275"
REGION = "ap-south-1"


def main() -> None:
    public_ip = EC2_IP_FILE.read_text(encoding="utf-8").strip() if EC2_IP_FILE.is_file() else "13.204.86.145"
    app_url = f"http://{public_ip}/"

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    t = doc.add_heading(
        "Deploying an Image Classification Application on AWS Using Docker, DVC, and CI/CD",
        level=0,
    )
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta = doc.add_paragraph(
        f"Student / AWS Account: {ACCOUNT_NAME}\nAWS Account ID: {ACCOUNT_ID}\nRegion: {REGION}\nDate: October 2026"
    )
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sections = [
        (
            "1. Introduction",
            "This project implements a modular image classification pipeline with DVC, containerizes the inference "
            "service with Docker, automates testing and image build using GitHub Actions (ci-cd.yaml), stores pipeline "
            "artifacts on AWS S3 as the DVC remote, and deploys the application on an AWS EC2 instance.",
        ),
        (
            "2. DVC Pipeline (dvc.yaml)",
            "Stages: data_collection -> data_processing -> feature_selection -> model_training -> model_evaluation. "
            "Hyperparameters are defined in params.yaml. Running dvc repro executes the full pipeline and tracks outputs "
            "in dvc.lock. Large artifacts are pushed to S3 remote storage.",
        ),
        (
            "3. AWS S3 DVC Remote Configuration",
            f"Bucket: s3://{BUCKET}/dvc-store\n"
            f"Account ID: {ACCOUNT_ID}\n"
            "Commands used:\n"
            "  dvc remote add -d storage s3://<bucket>/dvc-store\n"
            "  dvc remote modify storage region ap-south-1\n"
            "  dvc push",
        ),
        (
            "4. GitHub Actions CI/CD (ci-cd.yaml)",
            "Workflow jobs:\n"
            "  1) Run automated tests (python test.py) on push/PR\n"
            "  2) Build Docker image after tests pass on push to master/main\n"
            "Files: ci-cd.yaml (project root) and .github/workflows/ci-cd.yaml",
        ),
        (
            "5. Dockerfile",
            "Base image python:3.11-slim, installs requirements.txt + requirements-web.txt (Flask, Gunicorn), "
            "copies application source, exposes port 8000, runs gunicorn app.server:app.",
        ),
        (
            "6. EC2 Deployment",
            f"EC2 instance launched in {REGION} using deploy/launch_ec2.ps1 and deploy/ec2_user_data.sh.\n"
            "User data installs Docker, clones the GitHub repository, runs dvc pull from S3, builds the Docker image, "
            "and runs the container mapping host port 80 to container port 8000.",
        ),
        (
            "7. Verification",
            f"Public application URL: {app_url}\n"
            f"Health check: {app_url}health\n"
            "Upload an image on the home page or POST /predict with multipart form field 'image'.",
        ),
        (
            "8. Challenges and Solutions",
            "Challenge: Large DVC artifacts not suitable for Git.\n"
            "Solution: Configure S3 remote and use dvc push / dvc pull on EC2.\n\n"
            "Challenge: TensorFlow optional in CI; pipeline uses sklearn fallback when TF missing.\n"
            "Solution: test.py runs fast synthetic end-to-end pipeline.\n\n"
            "Challenge: EC2 bootstrap time for Docker build and dvc pull.\n"
            "Solution: Wait 5-10 minutes after instance launch before testing URL.",
        ),
        (
            "9. Conclusion",
            "The workflow demonstrates MLOps practices: reproducible DVC pipeline, cloud artifact storage, "
            "automated CI/CD, containerized deployment, and public verification on AWS EC2.",
        ),
        (
            "10. GitHub Repository",
            REPO,
        ),
    ]

    for title, text in sections:
        doc.add_heading(title, level=1)
        doc.add_paragraph(text)

    doc.add_heading("Screenshot Checklist (attach to submission)", level=1)
    checklist = [
        "dvc repro / dvc dag",
        "dvc push to S3 (AWS console bucket view)",
        "GitHub Actions ci-cd workflow green run",
        "docker build output",
        "EC2 instance running (AWS console)",
        "Browser access to public IP / health JSON",
    ]
    for item in checklist:
        doc.add_paragraph(item, style="List Bullet")

    # Embed CI screenshot if available
    ci_shot = ROOT / "submission" / "actions_latest_ci_green.png"
    if ci_shot.is_file():
        doc.add_paragraph("")
        doc.add_picture(str(ci_shot), width=Inches(5.8))
        cap = doc.add_paragraph("Figure: GitHub Actions CI successful run")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(f"Created: {OUT}")


if __name__ == "__main__":
    main()
